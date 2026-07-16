from pathlib import Path
from app.core.database import SessionLocal
from app.repositories.user_repository import UserRepository
from app.repositories.document_repository import DocumentRepository
from app.services.extraction.extractor_factory import ExtractorFactory
from docx import Document as DocxDocument

db = SessionLocal()
user = UserRepository(db).get_by_email("user@example.com")
docs = DocumentRepository(db).list_by_owner(user.id)

target = next((d for d in docs if "MHK-RGUKTN" in d.filename), None)
if not target:
    print("Document not found - check filename match")
    exit(1)

print(f"File: {target.filename}")
print(f"Storage path: {target.storage_path}\n")

extractor = ExtractorFactory.get_extractor(target.file_type)
chunks = extractor.extract(Path(target.storage_path))

print(f"Total extracted segments (before chunking): {len(chunks)}\n")
for i, c in enumerate(chunks, 1):
    print(f"[{i}] metadata={c.metadata}")
    print(f"    content: {c.content[:150]}")
    print()



doc = DocxDocument(target.storage_path)
print(f"Number of tables in document: {len(doc.tables)}\n")

for t_idx, table in enumerate(doc.tables):
    print(f"--- Table {t_idx} ({len(table.rows)} rows) ---")
    for r_idx, row in enumerate(table.rows):
        cells = [c.text.strip() for c in row.cells]
        print(f"  Row {r_idx}: {cells}")
    print()

db.close()