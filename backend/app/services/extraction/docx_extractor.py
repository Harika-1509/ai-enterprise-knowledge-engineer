from pathlib import Path

from docx import Document as DocxDocument

from app.services.extraction.base_extractor import BaseExtractor, ExtractedChunk


class DOCXExtractor(BaseExtractor):
    def extract(self, file_path: Path) -> list[ExtractedChunk]:
        doc = DocxDocument(str(file_path))
        chunks: list[ExtractedChunk] = []

        for i, paragraph in enumerate(doc.paragraphs):
            text = paragraph.text.strip()

            # Skip empty paragraphs
            if not text:
                continue

            # Skip heading-only paragraphs such as:
            # "Note:"
            # "Eligibility Criterion:"
            # "Selection Process:"
            # These contain almost no semantic information and
            # create noisy embeddings.
            words = text.split()
            if (
                len(words) <= 3
                and text.endswith(":")
            ):
                continue

            chunks.append(
                ExtractedChunk(
                    content=text,
                    metadata={
                        "paragraph_index": i,
                        "source_type": "docx",
                    },
                )
            )

        for t_idx, table in enumerate(doc.tables):
            chunks.extend(self._extract_table(table, t_idx))

        return chunks

    def _extract_table(self, table, t_idx: int) -> list[ExtractedChunk]:
        rows = table.rows
        if not rows:
            return []

        header = [cell.text.strip() for cell in rows[0].cells]
        non_empty_header_cells = [h for h in header if h]

        # A real multi-column header has more than one meaningful column
        # name (e.g. "Model", "Accuracy"). If only column 0 has text,
        # this is actually a key-value FORM table.
        is_form_table = len(non_empty_header_cells) <= 1

        if is_form_table:
            return self._extract_form_table(rows, t_idx)

        return self._extract_record_table(rows, header, t_idx)

    def _extract_form_table(self, rows, t_idx: int) -> list[ExtractedChunk]:
        chunks = []

        for r_idx, row in enumerate(rows):
            cells = [cell.text.strip() for cell in row.cells]

            if not cells or not cells[0]:
                continue

            label = cells[0]

            # Deduplicate repeated values while preserving order.
            seen = set()
            values = []

            for v in cells[1:]:
                if v and v != label and v not in seen:
                    values.append(v)
                    seen.add(v)

            if not values:
                continue

            value_text = "; ".join(values)
            content = f"{label}: {value_text}."

            chunks.append(
                ExtractedChunk(
                    content=content,
                    metadata={
                        "table_index": t_idx,
                        "row_index": r_idx,
                        "source_type": "docx_form_table",
                    },
                )
            )

        return chunks

    def _extract_record_table(
        self,
        rows,
        header: list[str],
        t_idx: int,
    ) -> list[ExtractedChunk]:

        chunks = []

        for r_idx, row in enumerate(rows[1:], start=1):
            values = [cell.text.strip() for cell in row.cells]

            pairs = [(h, v) for h, v in zip(header, values) if h and v]

            if not pairs:
                continue

            row_text = ", ".join(f"{h} is {v}" for h, v in pairs) + "."

            chunks.append(
                ExtractedChunk(
                    content=row_text,
                    metadata={
                        "table_index": t_idx,
                        "row_index": r_idx,
                        "source_type": "docx_table",
                    },
                )
            )

        return chunks