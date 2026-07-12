from pathlib import Path

from pptx import Presentation

from app.services.extraction.base_extractor import BaseExtractor, ExtractedChunk


class PPTXExtractor(BaseExtractor):
    def extract(self, file_path: Path) -> list[ExtractedChunk]:
        prs = Presentation(str(file_path))
        chunks: list[ExtractedChunk] = []

        for slide_number, slide in enumerate(prs.slides, start=1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = "".join(run.text for run in paragraph.runs).strip()
                        if text:
                            texts.append(text)

            slide_text = "\n".join(texts)
            if slide_text:
                chunks.append(
                    ExtractedChunk(
                        content=slide_text,
                        metadata={"slide_number": slide_number, "source_type": "pptx"},
                    )
                )

            # Speaker notes are often the most information-dense part of a deck
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    chunks.append(
                        ExtractedChunk(
                            content=notes,
                            metadata={
                                "slide_number": slide_number,
                                "source_type": "pptx_notes",
                            },
                        )
                    )
        return chunks