from pathlib import Path

from openpyxl import load_workbook

from app.services.extraction.base_extractor import BaseExtractor, ExtractedChunk


class XLSXExtractor(BaseExtractor):
    def extract(self, file_path: Path) -> list[ExtractedChunk]:
        workbook = load_workbook(str(file_path), data_only=True)
        chunks: list[ExtractedChunk] = []

        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            header = None

            for row_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = [str(cell) if cell is not None else "" for cell in row]
                if not any(values):
                    continue

                if header is None:
                    header = values
                    continue

                # Pair header with row values so the text is self-describing,
                # e.g. "Revenue: 50000 | Region: APAC" instead of a bare CSV row.
                row_text = " | ".join(
                    f"{h}: {v}" for h, v in zip(header, values) if h
                )
                if row_text:
                    chunks.append(
                        ExtractedChunk(
                            content=row_text,
                            metadata={
                                "sheet_name": sheet_name,
                                "row_number": row_idx,
                                "source_type": "xlsx",
                            },
                        )
                    )
        return chunks